#!/usr/bin/env python3
"""
End-to-end check that the v1.36 security context findings actually reach the Excel report.

The reporting methods swallow KeyError to stay quiet when a check finds nothing, which would also
hide a genuine column or sheet-name mistake, so these tests assert on the written workbook.
"""

import json
import os
import sys

import pandas as pd
import pytest
from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from kubenumerate import Kubenumerate  # noqa: E402


@pytest.fixture
def workbook(tmp_path):
    """Audit a pod exercising every new check, write the report, and return the workbook."""
    out_dir = tmp_path / "kubectl"
    out_dir.mkdir()
    items = [
        {
            "kind": "Pod",
            "apiVersion": "v1",
            "metadata": {"name": "unmasked-pod", "namespace": "prod"},
            "spec": {
                "hostUsers": False,
                "containers": [{"name": "builder", "securityContext": {"procMount": "Unmasked"}}],
            },
        },
        {
            "kind": "Pod",
            "apiVersion": "v1",
            "metadata": {"name": "anomalous-pod", "namespace": "prod"},
            "spec": {"containers": [{"name": "sneaky", "securityContext": {"procMount": "Unmasked"}}]},
        },
        {
            "kind": "Pod",
            "apiVersion": "v1",
            "metadata": {"name": "plain-pod", "namespace": "prod"},
            "spec": {"containers": [{"name": "app"}]},
        },
    ]
    with open(out_dir / "pods.json", "w") as f:
        json.dump({"items": items}, f)

    excel_file = tmp_path / "report.xlsx"
    kube = Kubenumerate(
        cluster_version="1.36.0",
        dry_run=False,
        kubectl_output_path=str(out_dir),
        excel_file=str(excel_file),
        verbosity=0,
    )
    df = kube.generate_kubeaudit_equivalent_df_from_kubectl()
    with pd.ExcelWriter(str(excel_file), engine="xlsxwriter", mode="w") as writer:
        kube.proc_mount(df, writer)
        kube.user_namespace(df, writer)
    return load_workbook(str(excel_file))


def sheet_text(workbook, sheet_name):
    return "\n".join(" ".join(str(cell) for cell in row if cell is not None) for row in workbook[sheet_name].values)


def test_proc_mount_unmasked_sheet_is_written(workbook):
    assert "Proc Mount - Unmasked" in workbook.sheetnames
    text = sheet_text(workbook, "Proc Mount - Unmasked")
    assert "builder" in text
    assert "sneaky" in text


def test_proc_mount_anomaly_sheet_only_lists_the_unprotected_pod(workbook):
    assert "Proc Mount - No UserNS" in workbook.sheetnames
    text = sheet_text(workbook, "Proc Mount - No UserNS")
    assert "anomalous-pod" in text
    assert "unmasked-pod" not in text


def test_proc_mount_sheet_reports_user_namespace_mitigation(workbook):
    text = sheet_text(workbook, "Proc Mount - Unmasked")
    assert "Mitigated by User Namespace" in text


def test_user_namespace_sheet_is_written(workbook):
    assert "User Namespace - Not Enabled" in workbook.sheetnames
    text = sheet_text(workbook, "User Namespace - Not Enabled")
    assert "plain-pod" in text
    assert "anomalous-pod" in text
    assert "unmasked-pod" not in text


def test_all_sheet_names_within_excel_limit(workbook):
    for sheet_name in workbook.sheetnames:
        assert len(sheet_name) <= 31, f"{sheet_name} exceeds Excel's 31-character sheet name limit"
